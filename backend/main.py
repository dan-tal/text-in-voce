import json
import re
import subprocess
import uuid
import urllib.request
import wave
from collections import Counter
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
import docx
import pdfplumber
from piper import PiperVoice

BASE_DIR = Path(__file__).parent
MODELS_DIR = BASE_DIR / "models"
STATIC_DIR = BASE_DIR / "static"
AUDIO_DIR = STATIC_DIR / "audio"
FRONTEND_DIR = BASE_DIR / "frontend"

MODEL_PATH = MODELS_DIR / "ro_RO-mihai-medium.onnx"
CONFIG_PATH = MODELS_DIR / "ro_RO-mihai-medium.onnx.json"

MODEL_URL = "https://huggingface.co/rhasspy/piper-voices/resolve/main/ro/ro_RO/mihai/medium/ro_RO-mihai-medium.onnx"
CONFIG_URL = MODEL_URL + ".json"

AUDIO_DIR.mkdir(parents=True, exist_ok=True)
MODELS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title="Text în voce")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

_voice = None


def _download_if_missing() -> None:
    if not MODEL_PATH.exists():
        urllib.request.urlretrieve(MODEL_URL, MODEL_PATH)
    if not CONFIG_PATH.exists():
        urllib.request.urlretrieve(CONFIG_URL, CONFIG_PATH)


def get_voice() -> PiperVoice:
    global _voice
    if _voice is None:
        _download_if_missing()
        _voice = PiperVoice.load(str(MODEL_PATH), config_path=str(CONFIG_PATH))
    return _voice


@app.on_event("startup")
def startup() -> None:
    # Ensure the model is present (and loaded) before the first request.
    get_voice()


SENTENCE_RE = re.compile(r"(?<=[.!?])\s+(?=[A-ZĂÂÎȘȚ0-9\"„])")


STYLE_KEYS = ("hl", "b", "i", "u")


def _run(text: str, **style) -> dict:
    return {"t": text, **{k: v for k, v in style.items() if v and k in STYLE_KEYS}}


def _style_of(run: dict) -> tuple:
    return tuple(run.get(k) for k in STYLE_KEYS)


def split_styled_sentences(runs: list[dict]) -> list[tuple[str, list[dict]]]:
    """Split a block (list of styled runs) into sentences, keeping per-char styles.

    Returns [(sentence_text, runs)] where runs is [] if nothing is styled.
    """
    chars: list[tuple[str, tuple]] = []
    for run in runs:
        style = _style_of(run)
        for ch in run["t"]:
            if ch.isspace():
                if chars and chars[-1][0] == " ":
                    continue
                chars.append((" ", style))
            else:
                chars.append((ch, style))
    while chars and chars[0][0] == " ":
        chars.pop(0)
    while chars and chars[-1][0] == " ":
        chars.pop()
    if not chars:
        return []

    text = "".join(c for c, _ in chars)
    out = []
    start = 0
    bounds = [(m.start(), m.end()) for m in SENTENCE_RE.finditer(text)]
    for b_start, b_end in bounds + [(len(text), len(text))]:
        seg = chars[start:b_start]
        if seg:
            grouped: list[dict] = []
            prev = None
            for ch, style in seg:
                if style == prev:
                    grouped[-1]["t"] += ch
                else:
                    grouped.append({"t": ch, **{k: v for k, v in zip(STYLE_KEYS, style) if v}})
                    prev = style
            styled = any(_style_of(r) != (None,) * len(STYLE_KEYS) for r in grouped)
            out.append(("".join(c for c, _ in seg), grouped if styled else []))
        start = b_end
    return out


def _extract_blocks_txt(text: str) -> list[dict]:
    blocks = []
    for para in re.split(r"\n\s*\n", text):
        joined = " ".join(line.strip() for line in para.splitlines() if line.strip()).strip()
        if joined:
            blocks.append({"type": "paragraph", "runs": [_run(joined)]})
    return blocks


HIGHLIGHT_COLORS = {
    "YELLOW": "#ffff00",
    "BRIGHT_GREEN": "#00ff00",
    "TURQUOISE": "#00ffff",
    "PINK": "#ff00ff",
    "BLUE": "#0000ff",
    "RED": "#ff0000",
    "DARK_BLUE": "#000080",
    "TEAL": "#008080",
    "GREEN": "#008000",
    "VIOLET": "#800080",
    "DARK_RED": "#800000",
    "DARK_YELLOW": "#808000",
    "GRAY_50": "#808080",
    "GRAY_25": "#c0c0c0",
    "BLACK": "#000000",
}


def _docx_run_highlight(run) -> str | None:
    color = run.font.highlight_color
    if color is not None:
        name = getattr(color, "name", str(color)).upper()
        return HIGHLIGHT_COLORS.get(name, "#ffff00")
    shd = run._r.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}shd")
    if shd is not None:
        fill = shd.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill")
        if fill and re.fullmatch(r"[0-9A-Fa-f]{6}", fill) and fill.lower() != "ffffff":
            return "#" + fill.lower()
    return None


def _docx_paragraph_runs(p) -> list[dict]:
    runs = [
        _run(r.text, hl=_docx_run_highlight(r), b=r.bold, i=r.italic, u=bool(r.underline))
        for r in p.runs
        if r.text
    ]
    # Runs inside hyperlinks etc. are not in p.runs; don't lose text if so.
    if "".join(r["t"] for r in runs).strip() != p.text.strip():
        return [_run(p.text)]
    return runs


def _extract_blocks_docx(content: bytes) -> list[dict]:
    tmp_path = BASE_DIR / f"tmp_{uuid.uuid4().hex}.docx"
    tmp_path.write_bytes(content)
    try:
        document = docx.Document(str(tmp_path))
        blocks = []
        for p in document.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            runs = _docx_paragraph_runs(p)
            style_name = (p.style.name or "").lower() if p.style else ""
            if "heading" in style_name or "title" in style_name:
                match = re.search(r"(\d+)", style_name)
                level = min(int(match.group(1)), 6) if match else 1
                blocks.append({"type": "heading", "level": level, "runs": runs})
            else:
                blocks.append({"type": "paragraph", "runs": runs})
        for table in document.tables:
            for row in table.rows:
                cells_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if cells_text:
                    blocks.append({"type": "paragraph", "runs": [_run(cells_text)]})
        return blocks
    finally:
        tmp_path.unlink(missing_ok=True)


def _pdf_fill_hex(color) -> str | None:
    """Hex for a saturated (non white/gray/black) fill colour, else None."""
    if not isinstance(color, (tuple, list)):
        return None
    if len(color) == 3:
        r, g, b = color
    elif len(color) == 4:
        c, m, y, k = color
        r, g, b = (1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)
    else:
        return None
    if max(r, g, b) - min(r, g, b) < 0.3:
        return None
    return "#%02x%02x%02x" % (round(r * 255), round(g * 255), round(b * 255))


def _pdf_highlight_rects(page) -> list[tuple[float, float, float, float, str]]:
    rects = []
    for r in page.rects:
        if not r.get("fill"):
            continue
        hex_color = _pdf_fill_hex(r.get("non_stroking_color"))
        # Skip big filled areas (backgrounds, table cells), keep text-sized marks.
        if hex_color and r["height"] < 40 and r["width"] < page.width * 0.95:
            rects.append((r["x0"], r["top"], r["x1"], r["bottom"], hex_color))
    return rects


def _pdf_word_run(w: dict, rects) -> dict:
    font = (w.get("fontname") or "").lower()
    cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
    hl = next((c for x0, t, x1, bt, c in rects if x0 - 1 <= cx <= x1 + 1 and t - 1 <= cy <= bt + 1), None)
    return _run(
        w["text"],
        hl=hl,
        b="bold" in font or "black" in font,
        i="italic" in font or "oblique" in font,
    )


def _extract_blocks_pdf(content: bytes) -> list[dict]:
    tmp_path = BASE_DIR / f"tmp_{uuid.uuid4().hex}.pdf"
    tmp_path.write_bytes(content)
    try:
        blocks: list[dict] = []
        pages_lines = []
        all_sizes = []

        with pdfplumber.open(str(tmp_path)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                words = page.extract_words(extra_attrs=["size", "fontname"])
                rects = _pdf_highlight_rects(page)
                if not words:
                    continue
                lines_by_top: dict[float, list] = {}
                for w in words:
                    key = round(w["top"])
                    lines_by_top.setdefault(key, []).append(w)
                page_lines = []
                for top in sorted(lines_by_top):
                    line_words = sorted(lines_by_top[top], key=lambda w: w["x0"])
                    text = " ".join(w["text"] for w in line_words)
                    line_runs = []
                    for w in line_words:
                        if line_runs:
                            line_runs.append(_run(" "))
                        line_runs.append(_pdf_word_run(w, rects))
                    avg_size = sum(w["size"] for w in line_words) / len(line_words)
                    page_lines.append({"top": top, "text": text, "runs": line_runs, "size": avg_size})
                    all_sizes.append(avg_size)
                if page_lines:
                    pages_lines.append((page_no, page_lines))

        if not all_sizes:
            return []

        body_size = Counter(round(s) for s in all_sizes).most_common(1)[0][0]

        for page_no, page_lines in pages_lines:
            body_gaps = [
                page_lines[i + 1]["top"] - page_lines[i]["top"]
                for i in range(len(page_lines) - 1)
                if page_lines[i]["size"] <= body_size * 1.15
                and page_lines[i + 1]["size"] <= body_size * 1.15
            ]
            # Regular single-line spacing is the smallest recurring gap; paragraph
            # breaks show up as outliers well above it. Median gets skewed with
            # few lines, so anchor on the minimum instead.
            normal_gap = min(body_gaps) if body_gaps else 15

            current_lines: list[list[dict]] = []
            current_is_heading = None
            prev_top = None

            def flush():
                if current_lines:
                    runs = []
                    for line_runs in current_lines:
                        if runs:
                            runs.append(_run(" "))
                        runs.extend(line_runs)
                    if "".join(r["t"] for r in runs).strip():
                        block = {"type": "paragraph", "runs": runs, "page": page_no}
                        if current_is_heading:
                            block.update(type="heading", level=2)
                        blocks.append(block)

            for line in page_lines:
                is_heading = line["size"] > body_size * 1.15
                gap = (line["top"] - prev_top) if prev_top is not None else 0
                starts_new_block = (
                    prev_top is None
                    or is_heading != current_is_heading
                    or gap > normal_gap * 1.6
                )
                if starts_new_block:
                    flush()
                    current_lines = [line["runs"]]
                    current_is_heading = is_heading
                else:
                    current_lines.append(line["runs"])
                prev_top = line["top"]
            flush()

        return blocks
    finally:
        tmp_path.unlink(missing_ok=True)


def extract_blocks(filename: str, content: bytes) -> list[dict]:
    suffix = Path(filename).suffix.lower()

    if suffix == ".txt":
        return _extract_blocks_txt(content.decode("utf-8", errors="ignore"))
    if suffix == ".docx":
        return _extract_blocks_docx(content)
    if suffix == ".pdf":
        return _extract_blocks_pdf(content)

    raise HTTPException(400, f"Format neacceptat: {suffix}. Folosește .docx, .pdf sau .txt.")


@app.post("/api/upload")
async def upload(file: UploadFile = File(...)):
    content = await file.read()
    raw_blocks = extract_blocks(file.filename or "", content)
    is_pdf = Path(file.filename or "").suffix.lower() == ".pdf"
    if not raw_blocks:
        raise HTTPException(400, "Nu s-a găsit text în fișier.")

    voice = get_voice()
    session_id = uuid.uuid4().hex[:12]
    session_dir = AUDIO_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    sentences = []
    blocks = []
    global_index = 0

    for raw_block in raw_blocks:
        block_sentences = split_styled_sentences(raw_block["runs"])
        if not block_sentences:
            continue

        indices = []
        for sentence, sentence_runs in block_sentences:
            wav_path = session_dir / f"{global_index}.wav"
            with wave.open(str(wav_path), "wb") as wav_file:
                voice.synthesize(sentence, wav_file)
            with wave.open(str(wav_path), "rb") as wav_file:
                frames = wav_file.getnframes()
                rate = wav_file.getframerate()
                duration = frames / float(rate) if rate else 0.0
            entry = {
                "index": global_index,
                "text": sentence,
                "audio_url": f"/audio/{session_id}/{global_index}.wav",
                "duration": duration,
            }
            if sentence_runs:
                entry["runs"] = sentence_runs
            if raw_block.get("page"):
                entry["page"] = raw_block["page"]
            sentences.append(entry)
            indices.append(global_index)
            global_index += 1

        block_out = {"type": raw_block["type"], "sentence_indices": indices}
        if raw_block.get("level"):
            block_out["level"] = raw_block["level"]
        blocks.append(block_out)

    if not sentences:
        raise HTTPException(400, "Nu s-a găsit text în fișier.")

    mp3_url = _build_mp3(session_dir, len(sentences))

    original_url = None
    if is_pdf:
        (session_dir / "original.pdf").write_bytes(content)
        original_url = f"/audio/{session_id}/original.pdf"

    result = {
        "session_id": session_id,
        "sentences": sentences,
        "blocks": blocks,
        "mp3_url": mp3_url,
        "original_url": original_url,
    }
    (session_dir / "meta.json").write_text(json.dumps(result), encoding="utf-8")
    return result


SESSION_ID_RE = re.compile(r"^[0-9a-f]{6,32}$")


@app.get("/api/session/{session_id}")
async def get_session(session_id: str):
    if not SESSION_ID_RE.match(session_id):
        raise HTTPException(404, "Sesiune negăsită sau expirată.")
    meta_path = AUDIO_DIR / session_id / "meta.json"
    if not meta_path.exists():
        raise HTTPException(404, "Sesiune negăsită sau expirată.")
    return json.loads(meta_path.read_text(encoding="utf-8"))


def _build_mp3(session_dir: Path, sentence_count: int) -> str | None:
    combined_path = session_dir / "full.wav"
    mp3_path = session_dir / "full.mp3"

    try:
        with wave.open(str(combined_path), "wb") as out_wav:
            for index in range(sentence_count):
                with wave.open(str(session_dir / f"{index}.wav"), "rb") as in_wav:
                    if index == 0:
                        out_wav.setparams(in_wav.getparams())
                    out_wav.writeframes(in_wav.readframes(in_wav.getnframes()))

        subprocess.run(
            ["ffmpeg", "-y", "-i", str(combined_path), "-codec:a", "libmp3lame", "-qscale:a", "2", str(mp3_path)],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except (OSError, subprocess.CalledProcessError):
        return None
    finally:
        combined_path.unlink(missing_ok=True)

    return f"/audio/{session_dir.name}/full.mp3"


app.mount("/audio", StaticFiles(directory=str(AUDIO_DIR)), name="audio")
app.mount("/", StaticFiles(directory=str(FRONTEND_DIR), html=True), name="frontend")
