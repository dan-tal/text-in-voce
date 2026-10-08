"""Exercise real HTTP routes, extraction and WAV output without downloading TTS."""
import asyncio
import io
import json
import sys
import threading
import types
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import docx
import httpx
import pytest

# Piper 1.2 is a Linux deployment dependency. The fake produces valid WAV files
# and blocks on demand so that concurrency is verified without timing guesses.
sys.modules.setdefault("piper", types.SimpleNamespace(PiperVoice=object))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
import main


class FakeVoice:
    def synthesize(self, text, wav_file):
        wav_file.setnchannels(1)
        wav_file.setsampwidth(2)
        wav_file.setframerate(8000)
        wav_file.writeframes(b"\0\0" * 800)


@pytest.fixture
def backend(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(main, "PROCESSING_WORKERS", 2)
    monkeypatch.setattr(main, "get_voice", lambda: FakeVoice())
    monkeypatch.setattr(main, "_build_mp3", lambda *args: None)
    return tmp_path


async def wait_until(predicate):
    for _ in range(500):
        if predicate():
            return
        await asyncio.sleep(0.01)
    raise AssertionError("Worker did not reach the expected state")


def test_parallel_uploads_keep_http_responsive_and_sessions_isolated(backend, monkeypatch):
    release = threading.Event()
    lock = threading.Lock()
    started = []
    process_document = main.process_document

    def blocked(filename, content):
        with lock:
            started.append(filename)
        assert release.wait(5)
        return process_document(filename, content)

    monkeypatch.setattr(main, "process_document", blocked)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                requests = [asyncio.create_task(client.post("/api/upload", files={"file": (f"{i}.txt", f"Document {i}. Salut!")})) for i in range(3)]
                try:
                    await wait_until(lambda: len(started) == 2 and main.app.state.pending_documents == 3)
                    # Both workers are still blocked. HTTP must be served now.
                    response = await asyncio.wait_for(client.get("/"), timeout=1)
                    assert response.status_code == 200
                    assert len(started) == 2  # third document remains queued
                finally:
                    release.set()
                responses = await asyncio.gather(*requests)
                assert all(response.status_code == 200 for response in responses)
                sessions = [response.json() for response in responses]
                assert len({session["session_id"] for session in sessions}) == 3
                for i, session in enumerate(sessions):
                    assert session["sentences"][0]["text"] == f"Document {i}."
                    assert session["sentences"][0]["duration"] == 0.1
                    persisted = await client.get(f'/api/session/{session["session_id"]}')
                    assert persisted.json() == session
                    assert (backend / session["session_id"] / "0.wav").exists()
                assert main.app.state.pending_documents == 0

    asyncio.run(scenario())


def test_overload_and_cancelled_request_keep_capacity_reserved(backend, monkeypatch):
    release = threading.Event()
    started = threading.Event()
    monkeypatch.setattr(main, "MAX_PENDING_DOCUMENTS", 1)
    original = main.process_document

    def blocked(*args):
        started.set()
        assert release.wait(5)
        return original(*args)

    monkeypatch.setattr(main, "process_document", blocked)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                request = asyncio.create_task(client.post("/api/upload", files={"file": ("one.txt", "Salut.")}))
                try:
                    await wait_until(started.is_set)
                    request.cancel()
                    with pytest.raises(asyncio.CancelledError):
                        await request
                    response = await client.post("/api/upload", files={"file": ("two.txt", "Salut.")})
                    assert response.status_code == 429
                    assert response.headers["retry-after"] == "5"
                    assert main.app.state.pending_documents == 1
                finally:
                    release.set()
                await wait_until(lambda: main.app.state.pending_documents == 0)
                response = await client.post("/api/upload", files={"file": ("three.txt", "Salut.")})
                assert response.status_code == 200

    asyncio.run(scenario())


@pytest.mark.parametrize("filename,content,status", [
    ("file.exe", b"hello", 400), ("empty.txt", b" \n ", 400),
    ("bad.docx", b"corrupt", 400), ("bad.pdf", b"corrupt", 400),
    ("large.txt", b"x" * 101, 413),
])
def test_invalid_uploads_release_capacity(backend, monkeypatch, filename, content, status):
    monkeypatch.setattr(main, "MAX_UPLOAD_BYTES", 100)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                response = await client.post("/api/upload", files={"file": (filename, content)})
                assert response.status_code == status
                assert main.app.state.pending_documents == 0
                assert not list(backend.iterdir())

    asyncio.run(scenario())


def test_synthesis_failure_removes_partial_audio(backend, monkeypatch):
    class BrokenVoice(FakeVoice):
        def synthesize(self, text, wav_file):
            super().synthesize(text, wav_file)
            raise RuntimeError("private engine details")

    monkeypatch.setattr(main, "get_voice", lambda: BrokenVoice())
    with pytest.raises(main.HTTPException) as error:
        main.process_document("file.txt", b"Salut.")
    assert error.value.status_code == 500
    assert "private engine details" not in error.value.detail
    assert not list(backend.iterdir())


def test_worker_voices_are_reused_but_not_shared(monkeypatch):
    monkeypatch.setattr(main, "_voices", threading.local())
    monkeypatch.setattr(main, "_download_if_missing", lambda: None)
    monkeypatch.setattr(main, "PiperVoice", types.SimpleNamespace(load=lambda *args, **kwargs: object()))
    barrier = threading.Barrier(2)

    def get_twice():
        first = main.get_voice()
        barrier.wait(timeout=5)
        assert main.get_voice() is first
        return first

    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(get_twice)
        second = pool.submit(get_twice)
        assert first.result() is not second.result()


def test_failed_model_download_is_not_published(monkeypatch, tmp_path):
    monkeypatch.setattr(main, "MODEL_PATH", tmp_path / "model.onnx")
    monkeypatch.setattr(main, "CONFIG_PATH", tmp_path / "model.json")

    def interrupted(url, destination):
        destination.write_bytes(b"partial download")
        raise OSError("network failure")

    monkeypatch.setattr(main.urllib.request, "urlretrieve", interrupted)
    with pytest.raises(OSError):
        main._download_if_missing()
    assert not list(tmp_path.iterdir())


def test_bom_txt_and_docx_structure(backend):
    assert main.extract_blocks("file.txt", b"\xef\xbb\xbfSalut.")[0]["runs"][0]["t"] == "Salut."
    document = docx.Document()
    document.add_heading("Titlu", level=1)
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Tabel."
    document.add_paragraph("Salut.")
    content = io.BytesIO()
    document.save(content)
    result = main.process_document("file.docx", content.getvalue())
    assert result["blocks"][0] == {"type": "heading", "level": 1, "sentence_indices": [0]}
    assert result["sentences"][1]["text"] == "Tabel."
    assert result["sentences"][2]["text"] == "Salut."


def test_mp3_timeout_preserves_sentence_audio_and_removes_partial_export(backend, monkeypatch):
    result = main.process_document("file.txt", b"Salut.")
    session_dir = backend / result["session_id"]
    # Undo this fixture's MP3 stub to exercise the export failure path.
    monkeypatch.undo()

    def timeout(*args, **kwargs):
        (session_dir / "full.mp3").write_bytes(b"partial")
        raise main.subprocess.TimeoutExpired("ffmpeg", kwargs["timeout"])

    monkeypatch.setattr(main.subprocess, "run", timeout)
    assert main._build_mp3(session_dir, 1) is None
    assert (session_dir / "0.wav").exists()
    assert not (session_dir / "full.wav").exists()
    assert not (session_dir / "full.mp3").exists()


def make_pdf(page_texts):
    """Small valid fixture, including deliberately empty source pages."""
    count = len(page_texts)
    kids = " ".join(f"{4 + i * 2} 0 R" for i in range(count))
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               f"<< /Type /Pages /Kids [{kids}] /Count {count} >>".encode(),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"]
    for i, text in enumerate(page_texts):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.extend([
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents {5 + i * 2} 0 R >>".encode(),
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        ])
    output = b"%PDF-1.4\n"
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(output)
    output += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    output += b"".join(f"{offset:010} 00000 n \n".encode() for offset in offsets[1:])
    return output + f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()


def test_pdf_keeps_original_numbering_blank_pages_and_styles(backend):
    content = make_pdf(["First page.", "", "Third page.", ""])
    result = main.process_document("document.pdf", content)
    assert result["page_count"] == 4
    assert [sentence["page"] for sentence in result["sentences"]] == [1, 3]
    assert [block["page"] for block in result["blocks"]] == [1, 3]
    assert result["sentences"][0]["runs"][0]["b"] is True
    assert result["original_url"].endswith("/original.pdf")
    assert (backend / result["session_id"] / "original.pdf").read_bytes() == content


def test_docx_styles_survive_concurrent_processing_refactor(backend):
    from docx.enum.text import WD_COLOR_INDEX
    document = docx.Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Text marcat.")
    run.bold = True
    run.italic = True
    run.underline = True
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    content = io.BytesIO()
    document.save(content)
    result = main.process_document("styled.docx", content.getvalue())
    assert result["sentences"][0]["runs"] == [{"t": "Text marcat.", "hl": "#ffff00", "b": True, "i": True, "u": True}]


def test_shared_library_persists_renames_and_deletes_all_document_files(backend, monkeypatch):
    audio_mount = next(route.app for route in main.app.routes if getattr(route, "path", None) == "/audio")
    monkeypatch.setattr(audio_mount, "all_directories", [str(backend)])

    def export(directory, count):
        (directory / "full.mp3").write_bytes(b"test mp3")
        return f"/audio/{directory.name}/full.mp3"

    monkeypatch.setattr(main, "_build_mp3", export)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as first:
                response = await first.post("/api/upload", files={"file": ("Raport.pdf", make_pdf(["Salut."]))})
                assert response.status_code == 200
                session = response.json()
                session_id = session["session_id"]
                async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as second:
                    items = (await second.get("/api/library")).json()["documents"]
                    assert len(items) == 1 and items[0]["name"] == "Raport.pdf"
                    assert items[0]["sentence_count"] == 1 and items[0]["size_bytes"] > 0
                    assert "sentences" not in items[0]
                    assert (await second.patch(f"/api/library/{session_id}", json={"name": "  Raport nou  "})).json()["name"] == "Raport nou"
                # Reopening a DB connection/browser retains the title.
                assert (await first.get("/api/library")).json()["documents"][0]["name"] == "Raport nou"
                for name in ["0.wav", "full.mp3", "original.pdf"]:
                    assert (await first.get(f"/audio/{session_id}/{name}")).status_code == 200
                assert (await first.delete(f"/api/library/{session_id}")).status_code == 204
                assert not (backend / session_id).exists()
                assert (await first.get("/api/library")).json() == {"documents": []}
                assert (await first.get(f"/api/session/{session_id}")).status_code == 404
                assert (await first.get(f"/audio/{session_id}/0.wav")).status_code == 404
                assert (await first.delete(f"/api/library/{session_id}")).status_code == 404
                assert (await first.get("/audio/.library/catalogue.sqlite3")).status_code == 404
                assert (await first.get("/audio/%2elibrary/catalogue.sqlite3")).status_code == 404

    asyncio.run(scenario())


def test_library_imports_legacy_sessions_once_and_skips_incomplete_documents(backend):
    good = backend / "abcdef123456"
    good.mkdir()
    (good / "meta.json").write_text(json.dumps({"session_id": good.name, "sentences": [{"text": "Salut."}]}))
    broken = backend / "123456abcdef"
    broken.mkdir()
    (broken / "meta.json").write_text("incomplete")
    assert main.get_library()["documents"][0]["name"] == "Document abcdef123456"
    assert len(main.get_library()["documents"]) == 1
    assert main.rename_document(good.name, main.DocumentName(name="Document vechi"))["name"] == "Document vechi"
    assert main.get_library()["documents"][0]["name"] == "Document vechi"


def test_library_rejects_invalid_titles_and_paths(backend):
    session = main.process_document("file.txt", b"Salut.")
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
            path = f'/api/library/{session["session_id"]}'
            assert (await client.patch(path, json={"name": "   "})).status_code == 400
            assert (await client.patch(path, json={"name": "x" * 201})).status_code == 422
            assert (await client.patch("/api/library/invalid", json={"name": "Nume"})).status_code == 404
            assert (await client.delete("/api/library/invalid")).status_code == 404
            assert (await client.delete("/api/library/%2e%2e%2f.library")).status_code in {404, 405}
            assert (await client.get("/api/library")).json()["documents"][0]["name"] == "file.txt"
    asyncio.run(scenario())


def test_docx_highlight_none_remains_readable(backend):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document = docx.Document()
    run = document.add_paragraph().add_run("Text fără marker.")
    highlight = OxmlElement("w:highlight")
    highlight.set(qn("w:val"), "none")
    run._r.get_or_add_rPr().append(highlight)
    content = io.BytesIO()
    document.save(content)
    assert main.process_document("document.docx", content.getvalue())["sentences"][0]["text"] == "Text fără marker."


def test_library_finishes_interrupted_file_cleanup(backend):
    main.get_library()
    trash = backend / ".library" / ("deleted-" + "a" * 32)
    trash.mkdir()
    (trash / "leftover.wav").write_bytes(b"interrupted delete")
    assert main.get_library()["documents"] == []
    assert not trash.exists()


def test_audio_remarks_are_shared_served_validated_and_deleted_with_the_document(backend, monkeypatch):
    audio_mount = next(route.app for route in main.app.routes if getattr(route, "path", None) == "/audio")
    monkeypatch.setattr(audio_mount, "all_directories", [str(backend)])
    session = main.process_document("note.txt", b"Primul paragraf.\n\nAl doilea paragraf.")
    session_id = session["session_id"]
    path = f"/api/session/{session_id}/remarks"

    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
            assert (await client.get(path)).json() == {"remarks": []}
            created = await client.post(path, data={"sentence_index": "1", "duration": "4.5", "author": " Ana "},
                                        files={"audio": ("remarca.webm", b"opus-bytes", "audio/webm;codecs=opus")})
            assert created.status_code == 201
            remark = created.json()
            assert remark["sentence_index"] == 1 and remark["duration"] == 4.5 and remark["author"] == "Ana"
            assert remark["audio_url"].startswith(f"/audio/{session_id}/remarks/") and remark["audio_url"].endswith(".webm")
            # Another visitor sees it and can stream the audio.
            assert (await client.get(path)).json()["remarks"] == [remark]
            assert (await client.get(remark["audio_url"])).content == b"opus-bytes"

            bad = lambda **kw: client.post(path, **kw)
            wav = {"audio": ("x.wav", b"data", "audio/wav")}
            assert (await bad(data={"sentence_index": "0"}, files={"audio": ("x.html", b"<script>", "text/html")})).status_code == 415
            assert (await bad(data={"sentence_index": "0"}, files={"audio": ("x.webm", b"", "audio/webm")})).status_code == 400
            assert (await bad(data={"sentence_index": "99"}, files=wav)).status_code == 400
            assert (await bad(data={"sentence_index": "-1"}, files=wav)).status_code == 422
            limit = main.MAX_REMARK_BYTES
            monkeypatch.setattr(main, "MAX_REMARK_BYTES", 3)
            assert (await bad(data={"sentence_index": "0"}, files=wav)).status_code == 413
            monkeypatch.setattr(main, "MAX_REMARK_BYTES", limit)
            assert (await client.post("/api/session/aaaaaaaaaaaa/remarks", data={"sentence_index": "0"}, files=wav)).status_code == 404
            assert (await client.get("/api/session/invalid/remarks")).status_code == 404

            assert (await client.delete(f"{path}/{'0' * 12}")).status_code == 404
            assert (await client.delete(f"{path}/../x")).status_code in {404, 405}
            assert (await client.delete(f'{path}/{remark["id"]}')).status_code == 204
            assert (await client.get(path)).json() == {"remarks": []}
            assert not list((backend / session_id / "remarks").iterdir())

            kept = (await client.post(path, data={"sentence_index": "0"}, files=wav)).json()
            assert (await client.delete(f"/api/library/{session_id}")).status_code == 204
            assert (await client.get(kept["audio_url"])).status_code == 404
            assert (await client.get(path)).status_code == 404
            assert (await client.post(path, data={"sentence_index": "0"}, files=wav)).status_code == 404

    asyncio.run(scenario())


def test_audio_supports_range_requests_for_ios_safari_and_is_cacheable(backend, monkeypatch):
    # Safari refuses to play (or seek in) <audio> unless the server answers byte ranges with 206.
    audio_mount = next(route.app for route in main.app.routes if getattr(route, "path", None) == "/audio")
    monkeypatch.setattr(audio_mount, "all_directories", [str(backend)])
    session = main.process_document("note.txt", b"Primul paragraf.")
    wav = f'/audio/{session["session_id"]}/0.wav'
    total = (backend / session["session_id"] / "0.wav").stat().st_size

    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
            probe = await client.get(wav, headers={"Range": "bytes=0-1"})   # what Safari sends first
            assert probe.status_code == 206
            assert probe.headers["content-range"] == f"bytes 0-1/{total}"
            assert probe.headers["accept-ranges"] == "bytes"
            assert len(probe.content) == 2
            tail = await client.get(wav, headers={"Range": "bytes=44-"})
            assert tail.status_code == 206 and len(tail.content) == total - 44
            full = await client.get(wav)
            assert full.status_code == 200 and len(full.content) == total
            for response in (probe, full):
                assert response.headers["cache-control"] == "public, max-age=86400"
            # Mutable bookkeeping files are never cached by phones.
            meta = await client.get(f'/audio/{session["session_id"]}/meta.json')
            assert meta.status_code == 200 and "max-age" not in meta.headers.get("cache-control", "")
            assert (await client.get(f'/audio/{session["session_id"]}/.hidden')).status_code == 404

    asyncio.run(scenario())


def test_frontend_is_revalidated_and_installable_on_phones():
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
            # A phone that kept the tab open must never mix a new page with an old script.
            for path in ("/", "/script.js", "/style.css", "/manifest.webmanifest"):
                response = await client.get(path)
                assert response.status_code == 200, path
                assert response.headers["cache-control"] == "no-cache", path
                assert response.headers["etag"], path                        # so revalidation is a cheap 304
                revalidated = await client.get(path, headers={"If-None-Match": response.headers["etag"]})
                assert revalidated.status_code == 304, path
            manifest = await client.get("/manifest.webmanifest")
            assert manifest.headers["content-type"].startswith("application/manifest+json")
            icons = manifest.json()["icons"]
            assert {icon["sizes"] for icon in icons} >= {"192x192", "512x512"}
            for icon in icons:
                served = await client.get(icon["src"])
                assert served.status_code == 200 and served.headers["content-type"] == icon["type"], icon["src"]
            assert (await client.get("/icons/apple-touch-icon.png")).status_code == 200

    asyncio.run(scenario())


def test_jobs_acknowledge_before_synthesis_and_retries_do_not_duplicate(backend, monkeypatch):
    release = threading.Event()
    started = threading.Event()

    class SlowVoice(FakeVoice):
        def synthesize(self, text, wav_file):
            started.set()
            assert release.wait(5)
            super().synthesize(text, wav_file)

    monkeypatch.setattr(main, "get_voice", lambda: SlowVoice())
    monkeypatch.setattr(main, "MAX_PENDING_DOCUMENTS", 1)
    job_id = "a" * 32

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                try:
                    response = await asyncio.wait_for(client.post("/api/jobs", data={"request_id": job_id},
                                                     files={"file": ("one.txt", "Salut. Bună ziua!")}), 1)
                    assert response.status_code == 202
                    assert response.headers["location"] == f"/api/jobs/{job_id}"
                    await wait_until(started.is_set)
                    polled = await asyncio.wait_for(client.get(response.headers["location"]), 1)
                    assert polled.headers["cache-control"] == "no-store"
                    assert polled.json()["status"] == "processing"
                    assert polled.json()["total_sentences"] == 2
                    duplicate = await client.post("/api/jobs", data={"request_id": job_id},
                                                  files={"file": ("one.txt", "Salut. Bună ziua!")})
                    assert duplicate.json()["job_id"] == job_id
                    assert main.app.state.pending_documents == 1
                    assert (await client.post("/api/jobs", files={"file": ("two.txt", "Salut.")})).status_code == 429
                    assert (await client.get("/")).status_code == 200
                    assert (await client.get("/api/jobs/invalid")).status_code == 404
                    assert (await client.get(f"/audio/.jobs/{job_id}.json")).status_code == 404
                finally:
                    release.set()
                await wait_until(lambda: main.app.state.pending_documents == 0)
                done = (await client.get(f"/api/jobs/{job_id}")).json()
                assert done["status"] == "ready" and done["completed_sentences"] == 2
                session = (await client.get(f'/api/session/{done["session_id"]}')).json()
                assert len(session["sentences"]) == 2
                assert len(main.get_library()["documents"]) == 1
        # A tab/server restart can recover completed status without input again.
        async with main.app.router.lifespan_context(main.app):
            assert main.jobs.read(backend, job_id)["status"] == "ready"

    asyncio.run(scenario())


@pytest.mark.parametrize("filename,content", [("bad.docx", b"broken"), ("empty.txt", b" ")])
def test_failed_jobs_report_errors_and_release_capacity(backend, filename, content):
    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                response = await client.post("/api/jobs", files={"file": (filename, content)})
                assert response.status_code == 202
                await wait_until(lambda: main.app.state.pending_documents == 0)
                data = (await client.get(response.headers["location"])).json()
                assert data["status"] == "failed" and data["error_status"] == 400
                assert data["error"]
                assert not [path for path in backend.iterdir() if not path.name.startswith(".")]
    asyncio.run(scenario())


def test_missing_voice_does_not_prevent_startup_or_expose_engine_details(backend, monkeypatch):
    def unavailable():
        raise OSError("private network credentials")
    monkeypatch.setattr(main, "get_voice", unavailable)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                assert (await client.get("/")).status_code == 200
                response = await client.post("/api/jobs", files={"file": ("one.txt", "Salut.")})
                await wait_until(lambda: main.app.state.pending_documents == 0)
                data = (await client.get(response.headers["location"])).json()
                assert data["status"] == "failed" and data["error_status"] == 503
                assert "private" not in data["error"]
    asyncio.run(scenario())


def test_interrupted_jobs_are_marked_failed_and_old_status_records_are_pruned(backend):
    pending = main.jobs.create(backend, "one.txt")
    done = main.jobs.create(backend, "two.txt")
    main.jobs.update(backend, done["job_id"], status="ready")
    path = backend / ".jobs" / f'{done["job_id"]}.json'
    data = json.loads(path.read_text())
    data["updated_at"] = 0
    path.write_text(json.dumps(data))
    main.jobs.recover(backend)
    assert main.jobs.read(backend, pending["job_id"])["status"] == "failed"
    assert main.jobs.read(backend, done["job_id"]) is None


def test_deleted_document_cannot_be_recreated_by_a_late_remark(backend):
    result = main.process_document("one.txt", b"Salut.")
    directory = backend / result["session_id"]
    assert main.library.delete(backend, result["session_id"])
    with pytest.raises(FileNotFoundError):
        main.remarks.add(directory, result["session_id"], 0, b"voice", "webm", 1, "")
    assert not directory.exists()


def test_delete_succeeds_even_if_private_disk_cleanup_must_be_retried(backend, monkeypatch):
    result = main.process_document("one.txt", b"Salut.")
    original = main.library.shutil.rmtree
    monkeypatch.setattr(main.library.shutil, "rmtree", lambda *args, **kwargs: (_ for _ in ()).throw(OSError("busy")))
    assert main.library.delete(backend, result["session_id"])
    assert not (backend / result["session_id"]).exists()
    monkeypatch.setattr(main.library.shutil, "rmtree", original)
    assert main.get_library()["documents"] == []
    assert not list((backend / ".library").glob("deleted-*"))


def test_disconnect_during_status_write_still_enqueues_the_document(backend, monkeypatch):
    release = threading.Event()
    started = threading.Event()
    create = main.jobs.create
    job_id = "c" * 32

    def slow_create(*args):
        started.set()
        assert release.wait(5)
        return create(*args)

    monkeypatch.setattr(main.jobs, "create", slow_create)

    async def scenario():
        async with main.app.router.lifespan_context(main.app):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(main.app), base_url="http://test") as client:
                request = asyncio.create_task(client.post("/api/jobs", data={"request_id": job_id},
                                                         files={"file": ("one.txt", "Salut.")}))
                try:
                    await wait_until(started.is_set)
                    request.cancel()
                    # Let cancellation reach the await of the status write.
                    await asyncio.sleep(0)
                    assert main.app.state.pending_documents == 1
                finally:
                    release.set()
                with pytest.raises(asyncio.CancelledError):
                    await request
                await wait_until(lambda: main.app.state.pending_documents == 0)
                retry = await client.post("/api/jobs", data={"request_id": job_id},
                                          files={"file": ("one.txt", "Salut.")})
                assert retry.status_code == 202 and retry.json()["status"] == "ready"
                assert len(main.get_library()["documents"]) == 1

    asyncio.run(scenario())


@pytest.mark.parametrize("data", [[], {"status": "ready"},
                                  {"job_id": "d" * 32, "status": [], "updated_at": 0},
                                  {"job_id": "d" * 32, "status": "ready", "updated_at": "invalid"}])
def test_corrupt_status_records_do_not_break_recovery_or_polling(backend, data):
    directory = backend / ".jobs"
    directory.mkdir()
    (directory / f'{"d" * 32}.json').write_text(json.dumps(data))
    main.jobs.recover(backend)
    assert main.jobs.read(backend, "d" * 32) is None
